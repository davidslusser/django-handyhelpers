import io
from datetime import datetime, timedelta
from datetime import timezone as dt_timezone
from unittest.mock import patch
from urllib.parse import parse_qs, urlsplit

import pytest
from django import forms
from django.contrib.auth.models import AnonymousUser, Group
from django.core.exceptions import ImproperlyConfigured, ValidationError
from django.core.management import call_command
from django.core.management.base import CommandError
from django.http import HttpResponse
from django.test import RequestFactory, override_settings
from django.urls import resolve
from openpyxl import load_workbook
from rest_framework import serializers
from rest_framework.request import Request
from rest_framework.test import APIRequestFactory

from handyhelpers.forms import HtmxModelForm
from handyhelpers.middleware import RequireLoginMiddleware
from handyhelpers.mixins.model_mixins import NonEmptyStringMixin
from handyhelpers.mixins.serializer_mixins import ExcludeNullMixin
from handyhelpers.mixins.view_mixins import FilterByQueryParamsMixin, HtmxViewMixin
from handyhelpers.permissions import InAllGroups, InAnyGroup
from handyhelpers.querysets import count_by_hour, count_by_month, count_by_week
from handyhelpers.serializers import FkReadWriteField, NoModelFlexSerializer
from handyhelpers.templatetags.handyhelper_tags import (
    get_filtered_pagination_link,
    inlist,
    nbsp,
    or_blank,
    percentage,
)
from handyhelpers.templatetags.model_meta_tags import db_table
from handyhelpers.utils import InvalidFilter, optional_import, update_query
from handyhelpers.views.calendar import CalendarView
from handyhelpers.views.export import CsvExportView, ExcelExportView
from handyhelpers.views.gui import (
    HandyHelperAboutView,
    HandyHelperListView,
    HandyHelperPaginatedListView,
)
from handyhelpers.views.htmx import (
    BuildBootstrapModalView,
    CreateModelModalView,
    HtmxOptionMultiFilterView,
)
from handyhelpers.views.report import (
    build_annual_progress_chart,
    build_annual_trend_chart,
    get_colors,
    get_timestamps,
)
from tests.models import Category, Configuration, Item, RequiredName, User

pytestmark = pytest.mark.django_db


class ItemForm(HtmxModelForm):
    class Meta:
        model = Item
        fields = ["name", "category", "active", "number"]


def request(path="/", method="get", data=None, htmx=False):
    result = getattr(RequestFactory(), method)(
        path, data or {}, HTTP_HX_REQUEST="true" if htmx else "false"
    )
    result.user = AnonymousUser()
    result.resolver_match = resolve("/")
    return result


def filter_view(parameters):
    view = FilterByQueryParamsMixin()
    view.request = RequestFactory().get("/", parameters)
    view.queryset = Item.objects.all()
    return view


@pytest.mark.parametrize(
    "params",
    [{"unknown": "x"}, {"name__nonsense": "x"}, {"number": "NaN"}, {"created_at__gte": "invalid"}],
)
def test_invalid_filters_are_rejected(params):
    with pytest.raises(InvalidFilter):
        filter_view(params).filter_by_query_params()
    response = HandyHelperListView.as_view(queryset=Item.objects.all())(request("/", data=params))
    assert response.status_code == 400


def test_filters_control_params_and_encoding():
    Item.objects.create(name="a&b", number=1)
    assert (
        filter_view({"name": "a&b", "page": "2", "display": "table"})
        .filter_by_query_params()
        .count()
        == 1
    )
    assert parse_qs(urlsplit(update_query("/items/?old=x", {"name": "a&b", "space": "a b"})).query)[
        "name"
    ] == ["a&b"]
    params = parse_qs(get_filtered_pagination_link(3, "name=a%26b&page=1&other=ok"))
    assert params == {"page": ["3"], "name": ["a&b"], "other": ["ok"]}
    with pytest.raises(InvalidFilter):
        filter_view({"number__in": "not a number"}).filter_by_query_params()


def test_request_isolation_and_htmx_header():
    class Modal(BuildBootstrapModalView):
        data = {"nested": {"value": 1}}

    left, right = Modal(), Modal()
    left.data["nested"]["value"] = 2
    assert right.data["nested"]["value"] == 1
    view = HtmxViewMixin()
    view.request = request()
    assert not view.is_htmx()
    view.request = request(htmx=True)
    assert view.is_htmx()


def test_login_prevents_side_effects_and_prefix_exemptions():
    calls = []
    middleware = RequireLoginMiddleware(
        lambda req: calls.append(req.path) or HttpResponse("called")
    )
    assert middleware(request("/protected/")).status_code == 302
    assert calls == []
    assert middleware(request("/administrator/")).status_code == 302
    assert middleware(request("/admin/login/")).status_code == 200
    json_request = RequestFactory().post("/protected/", data={}, content_type="application/json")
    json_request.user = AnonymousUser()
    assert middleware(json_request).status_code == 401


@pytest.mark.parametrize("permission", [InAllGroups, InAnyGroup])
def test_permission_matrix(permission):
    admin = User.objects.create_superuser(username="admin", password="independent")
    member = User.objects.create_user(username="member")
    group = Group.objects.create(name="operators")
    member.groups.add(group)
    view = permission()
    view.permission_dict = {"GET": ["operators"], "POST": []}
    for user, method, expected in [
        (AnonymousUser(), "GET", False),
        (member, "GET", True),
        (admin, "GET", True),
        (admin, "DELETE", False),
        (AnonymousUser(), "POST", True),
    ]:
        req = request(method=method.lower())
        req.user = user
        assert view.has_permission(req) == expected


def test_list_forms_and_pagination_do_not_leak():
    for number in range(4):
        Item.objects.create(name=str(number))

    class List(HandyHelperPaginatedListView):
        queryset = Item.objects.order_by("pk")
        paginate_by = 2
        create_form_obj = ItemForm

    left, right = List(), List()
    for view, path in [(left, "/?page=-1"), (right, "/?page=99")]:
        view.setup(request(path))
        view.object_list = view.get_queryset()
    first = left.get_context_data()
    second = right.get_context_data()
    assert first["page_obj"].number == 2 and second["page_obj"].number == 2
    assert first["create_form"] is not second["create_form"]
    assert first["create_form"]["modal_name"] == "create"


def test_bound_fields_preserve_widgets_and_prefixes():
    category = Category.objects.create(name="Selected")
    form = ItemForm(
        prefix="entry",
        initial={"name": "Initial", "category": category.pk, "active": True, "number": 0},
    )
    html = form.as_bs5()
    assert 'name="entry-name"' in html and 'value="Initial"' in html
    assert "selected" in html and "checked" in html and 'value="0"' in html
    assert "disabled" in form.as_bs5_disabled()
    invalid = ItemForm(data={"entry-name": ""}, prefix="entry")
    assert not invalid.is_valid()
    assert "invalid-feedback" in invalid.as_bs5()

    class UploadForm(forms.Form):
        file = forms.FileField()
        secret = forms.CharField(widget=forms.HiddenInput)

    from handyhelpers.templatetags.form_tags import bootstrap_field

    assert 'type="file"' in bootstrap_field(UploadForm()["file"])


def test_modal_submission_preserves_errors_and_saves():
    view = CreateModelModalView.as_view(form=ItemForm)
    invalid = view(request(method="post", data={"name": ""}, htmx=True))
    assert invalid.status_code == 422 and b"This field is required" in invalid.content
    assert b'data-bs-dismiss="modal">Create' not in invalid.content
    valid = view(request(method="post", data={"name": "Created", "number": 0}, htmx=True))
    assert valid.status_code == 204 and "handyhelpers:created" in valid["HX-Trigger"]
    assert Item.objects.filter(name="Created").exists()
    assert view(request(method="post", data={"name": "Rejected", "number": 0})).status_code == 400


def test_multi_filter_invalid_post_has_no_unbound_variable_or_csrf_filter():
    class FilterForm(forms.Form):
        number = forms.IntegerField(required=False)

    class Multi(HtmxOptionMultiFilterView):
        queryset = Item.objects.all()
        filter_form = FilterForm
        htmx_table_template_name = "handyhelpers/table/auditlog.htm"

    assert (
        Multi.as_view()(request(method="post", data={"number": "bad"}, htmx=True)).status_code
        == 400
    )
    assert (
        Multi.as_view()(
            request(
                method="post", data={"number": "0", "csrfmiddlewaretoken": "ignored"}, htmx=True
            )
        ).status_code
        == 200
    )


def test_safe_redirects(client):
    response = client.post(
        "/handyhelpers/filter_list_view",
        {"name": "a&b", "csrfmiddlewaretoken": "ignore"},
        HTTP_REFERER="https://evil.example/",
    )
    assert response["Location"] == "/?name=a%26b"
    assert client.get("/handyhelpers/show_all_list_view").get("Location") == "/"


def test_fk_field_respects_queryset():
    allowed = Category.objects.create(name="allowed")
    hidden = Category.objects.create(name="hidden")
    field = FkReadWriteField(queryset=Category.objects.filter(pk=allowed.pk))
    assert field.to_internal_value(allowed.pk) == allowed
    with pytest.raises(serializers.ValidationError):
        field.to_internal_value(hidden.pk)


def test_serializer_selection_and_null_exclusion():
    class Example(ExcludeNullMixin, NoModelFlexSerializer):
        name = serializers.CharField()
        optional = serializers.CharField(allow_null=True)

    req = Request(APIRequestFactory().get("/", {"fields": "name"}))
    assert Example({"name": "x", "optional": None}, context={"request": req}).data == {"name": "x"}
    assert Example({"name": "x", "optional": None}).data == {"name": "x"}


def test_exports_are_valid_and_formula_safe():
    Item.objects.create(name="=1+1", number=-2)
    csv_response = CsvExportView.as_view(queryset=Item.objects.all())(request())
    assert b",-2\r\n" in csv_response.content
    assert b"'=1+1" in csv_response.content and csv_response["Content-Disposition"].endswith(
        '.csv"'
    )
    excel_response = ExcelExportView.as_view(queryset=Item.objects.all(), filename="old.xls")(
        request()
    )
    workbook = load_workbook(io.BytesIO(excel_response.content))
    headers = [cell.value for cell in workbook.active[1]]
    cell = workbook.active.cell(2, headers.index("name") + 1)
    assert cell.value == "=1+1" and cell.data_type == "s"
    assert excel_response["Content-Disposition"].endswith('.xlsx"')


def test_empty_exports():
    assert b"name" in CsvExportView.as_view(queryset=Item.objects.none())(request()).content
    assert (
        load_workbook(
            io.BytesIO(ExcelExportView.as_view(queryset=Item.objects.none())(request()).content)
        ).active.max_row
        == 1
    )


def test_calendar_configured_field_and_boundaries():
    Item.objects.create(
        name="Leap event", event_at=datetime(2024, 2, 29, 12, tzinfo=dt_timezone.utc)
    )
    req = request("/calendar/2024/2/")
    req.resolver_match = resolve("/calendar/2024/2/")
    response = CalendarView.as_view(
        event_model=Item, event_model_date_field="event_at", use_htmx=False
    )(req, year=2024, month=2)
    assert b"Leap event" in response.content and b"/calendar/2024/3/" in response.content
    assert CalendarView.as_view()(req, year=2024, month=13).status_code == 400
    assert CalendarView().get_next(2024, 12) == (2025, 1)
    assert CalendarView().get_previous(2024, 1) == (2023, 12)


def test_calendar_reports_and_portable_aggregation():
    now = datetime(2024, 3, 31, 12, tzinfo=dt_timezone.utc)
    items = [
        Item.objects.create(name=str(i), event_at=now - timedelta(days=i * 30)) for i in range(14)
    ]
    data = [
        {
            "title": "Items",
            "queryset": Item.objects.all(),
            "dt_field": "event_at",
            "list_view": "/items/?existing=yes",
        }
    ]
    with (
        patch("handyhelpers.views.report.local_now", return_value=now),
        patch("handyhelpers.querysets._now", return_value=now),
    ):
        labels, months, trend = build_annual_trend_chart(data)
        assert len(months) == 12 and months[-1].month == 3
        _, _, progress = build_annual_progress_chart(data)
        assert progress[0]["data"][-1] == len(items)
        assert sum(trend[0]["data"]) <= len(items)
        assert len(count_by_hour(Item.objects.all(), "event_at")) == 24
        assert len(count_by_week(Item.objects.all(), "event_at")) == 52
        assert len(count_by_month(Item.objects.all(), "event_at")) == 12
        assert count_by_hour(Item.objects.all(), "event_at")[0] == 1
    assert get_timestamps(now)[2].day == 29
    assert len(get_colors(100)) == 100


def test_model_utilities_and_tags():
    assert Configuration.load().pk == 1
    with pytest.raises(ValidationError):
        RequiredName(name="").save()
    assert isinstance(RequiredName(), NonEmptyStringMixin)
    assert not inlist("ad", "admin,user") and inlist("admin", "admin,user")
    assert or_blank(0) == 0 and percentage(0) == "0.00%"
    assert "&lt;script&gt;" in nbsp("<script> hello")
    assert db_table(Item) == "tests_item"


def test_account_commands_preserve_identity_and_credentials():
    output = io.StringIO()
    with pytest.raises(CommandError):
        call_command("add_superuser", stdout=output)
    call_command("add_user", username="new", password="first", stdout=output)
    user = User.objects.get(username="new")
    item = Item.objects.create(name="relation")
    call_command("add_user", username="new", password="second", override=True, stdout=output)
    assert User.objects.get(username="new").pk == user.pk
    call_command(
        "add_superuser",
        username="new",
        generate_password=True,
        override=True,
        group="admins",
        stdout=output,
    )
    user.refresh_from_db()
    assert user.is_superuser and user.groups.filter(name="admins").exists()
    assert not user.check_password("new")
    assert Item.objects.filter(pk=item.pk).exists()


def test_admin_generator_force_and_app_config(tmp_path):
    output = tmp_path / "admin.py"
    call_command("create_admin", "tests", output_file=str(output))
    assert "ItemAdmin" in output.read_text()
    compile(output.read_text(), str(output), "exec")
    with pytest.raises(CommandError):
        call_command("create_admin", "tests", output_file=str(output))
    call_command("create_admin", "tests", output_file=str(output), force=True)


def test_optional_import_error_is_actionable():
    with pytest.raises(ImproperlyConfigured, match="django-handyhelpers\\[excel\\]"):
        optional_import("nonexistent_hh_optional", "excel")


def test_template_override_is_resolved_at_request_time():
    with override_settings(BASE_TEMPLATE="handyhelpers/handyhelpers_with_sidebar.htm"):
        response = HandyHelperAboutView.as_view()(request())
        assert b"body_content_wrapper" in response.content
